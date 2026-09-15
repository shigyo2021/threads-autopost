"""投稿管理表（Excel）: 投稿ログとストックから、商品ごとの一覧表を作る

    py sheet.py    管理表を作り直す

drafts.py の各コマンドと、tool.py での投稿後にも自動で作り直される。
表は毎回データから作り直すので、Excel上で編集しても元データには反映されない。
"""

import glob
import json
import os
import platform
import sys
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import _BASE_DIR
from drafts import USED_DIR, _live_posts, is_ready, load_stock, post_url

# Gitの管理外（files/ の1つ上）に置き、エクスプローラーから開きやすくする
SHEET_PATH = os.path.join(os.path.dirname(_BASE_DIR), "投稿管理表.xlsx")

STATUS_POSTED = "投稿済み"
STATUS_QUEUED = "予約済み"
STATUS_PLANNED = "投稿予定"
STATUS_UNPLANNED = "予定日なし"
STATUS_NO_TEXT = "紹介文未作成"

TEXT_SOURCE_LABELS = {"claude_code": "Claude Code", "api": "API"}

STATUS_FILLS = {
    STATUS_POSTED: "E2EFDA",
    STATUS_QUEUED: "D9D2E9",
    STATUS_PLANNED: "DDEBF7",
    STATUS_UNPLANNED: "FFF2CC",
    STATUS_NO_TEXT: "FCE4D6",
}

# (見出し, 列幅)
COLUMNS = [
    ("ステータス", 12),
    ("投稿予定日", 12),
    ("投稿時間", 9),
    ("投稿日", 12),
    ("投稿済み", 9),
    ("紹介文", 8),
    ("Threadsリンク", 13),
    ("商品名", 50),
    ("価格", 9),
    ("ショップ", 22),
    ("レビュー", 16),
    ("商品ページ", 45),
    ("登録したURL", 30),
    ("本文", 60),
    ("返信文", 45),
    ("紹介文の作成", 13),
    ("商品コード", 26),
    ("投稿回数", 9),
]


def _to_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(value[:10]).date()
    except ValueError:
        return None


def _review_label(draft: dict) -> str:
    if not draft.get("review_count"):
        return ""
    return f"★{draft['review_average']}（{draft['review_count']:,}件）"


def _load_used_drafts() -> dict:
    used = {}
    for path in glob.glob(os.path.join(USED_DIR, "*.json")):
        with open(path, "r", encoding="utf-8") as f:
            draft = json.load(f)
        used[draft["item_code"]] = draft
    return used


def build_rows() -> list[dict]:
    """商品ごとに1行。投稿ログ・使用済み下書き・ストックを item_code でまとめる"""
    rows = {}
    used = _load_used_drafts()

    for post in _live_posts():
        code = post.get("item_code")
        if not code:
            continue
        row = rows.setdefault(code, {"post_count": 0})
        row["post_count"] += 1
        row["posted_at"] = max(row.get("posted_at") or "", post.get("timestamp", ""))
        draft = used.get(code, {})
        row["draft"] = {**post, **draft, "affiliate_url": draft.get("affiliate_url") or post.get("url", "")}
        # text_source がない古いログは、botやtool.pyのAPI生成で作ったもの
        row["text_source"] = TEXT_SOURCE_LABELS.get(post.get("text_source"), "API")

    for draft in load_stock():
        code = draft["item_code"]
        row = rows.setdefault(code, {"post_count": 0})
        # 投稿済みの商品を再びストックした場合も、ストック側の情報で予定を表示する
        row["draft"] = draft
        row["in_stock"] = True
        row["text_source"] = "Claude Code" if is_ready(draft) else ""

    result = []
    for code, row in rows.items():
        draft = row["draft"]
        posted = row["post_count"] > 0 and not row.get("in_stock")
        if posted:
            status = STATUS_POSTED
        elif draft.get("queued_at"):
            status = STATUS_QUEUED
        elif not is_ready(draft):
            status = STATUS_NO_TEXT
        elif draft.get("planned_date"):
            status = STATUS_PLANNED
        else:
            status = STATUS_UNPLANNED

        result.append({
            "ステータス": status,
            "投稿予定日": None if posted else _to_date(draft.get("planned_date")),
            "投稿時間": row.get("posted_at", "")[11:16] if posted else (draft.get("planned_time") or ""),
            "投稿日": _to_date(row.get("posted_at")),
            "投稿済み": "済" if row["post_count"] else "未",
            "紹介文": "済" if draft.get("post_text", "").strip() else "未",
            # SNSボタンで作ったThreads用リンクを返信に使う（使った）か。r10.to は SNSボタンのリンク
            "Threadsリンク": "済" if draft.get("reply_link") or "r10.to" in str(draft.get("url", "")) else "未",
            "商品名": draft.get("name", ""),
            "価格": draft.get("price"),
            "ショップ": draft.get("shop", ""),
            "レビュー": _review_label(draft),
            "商品ページ": post_url(draft),
            "登録したURL": draft.get("source_url", ""),
            "本文": draft.get("post_text", ""),
            "返信文": draft.get("reply_text", ""),
            "紹介文の作成": row.get("text_source", ""),
            "商品コード": code,
            "投稿回数": row["post_count"],
        })

    # 未投稿を予定日順に上へ、投稿済みは新しい順に下へ
    unposted = sorted(
        (r for r in result if r["ステータス"] != STATUS_POSTED),
        key=lambda r: (r["投稿予定日"] or date.max, r["投稿時間"] or "99:99", r["商品名"]),
    )
    posted = sorted(
        (r for r in result if r["ステータス"] == STATUS_POSTED),
        key=lambda r: r["投稿日"] or date.min,
        reverse=True,
    )
    return unposted + posted


def export_sheet(path: str = SHEET_PATH) -> str:
    rows = build_rows()

    wb = Workbook()
    ws = wb.active
    ws.title = "投稿管理"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="44546A")
    for col, (title, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width

    titles = [title for title, _ in COLUMNS]
    for r, row in enumerate(rows, 2):
        for c, title in enumerate(titles, 1):
            cell = ws.cell(row=r, column=c, value=row[title])
            cell.alignment = Alignment(vertical="top")
            if title in ("投稿予定日", "投稿日") and row[title]:
                cell.number_format = "yyyy/mm/dd"
            elif title == "価格" and row[title] is not None:
                cell.number_format = "¥#,##0"
            elif title in ("商品ページ", "登録したURL") and row[title]:
                cell.hyperlink = row[title]
                cell.style = "Hyperlink"
                cell.alignment = Alignment(vertical="top")
            elif title in ("投稿時間", "投稿済み", "紹介文", "Threadsリンク", "投稿回数"):
                cell.alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=STATUS_FILLS[row["ステータス"]])

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(rows) + 1, 2)}"

    _add_guide_sheet(wb, rows)

    # Excelで開いたままだと上書きできないので、一時ファイルに書いてから置き換える
    tmp_path = path + ".tmp.xlsx"
    wb.save(tmp_path)
    try:
        os.replace(tmp_path, path)
    except PermissionError:
        os.remove(tmp_path)
        raise PermissionError(f"管理表がExcelで開かれているため更新できません。閉じてから `py sheet.py` を実行してください: {path}")
    return path


def _add_guide_sheet(wb: Workbook, rows: list[dict]):
    ws = wb.create_sheet("説明")
    counts = {status: sum(r["ステータス"] == status for r in rows) for status in STATUS_FILLS}
    lines = [
        ("最終更新", datetime.now().strftime("%Y/%m/%d %H:%M")),
        ("", ""),
        ("件数", ""),
        *[(f"  {status}", counts[status]) for status in STATUS_FILLS],
        ("", ""),
        ("ステータス", ""),
        (f"  {STATUS_POSTED}", "Threadsに投稿した商品"),
        (f"  {STATUS_QUEUED}", "GitHub Actionsに予約済み（時間になると自動で投稿される）"),
        (f"  {STATUS_PLANNED}", "紹介文があり、投稿予定日が決まっている（ストック）"),
        (f"  {STATUS_UNPLANNED}", "紹介文はあるが予定日がない → py drafts.py plan"),
        (f"  {STATUS_NO_TEXT}", "商品は登録済みだが紹介文がまだない"),
        ("", ""),
        ("調べ方", ""),
        ("  投稿済みか確認", "Ctrl+F で商品名・ショップ名・短縮URL（r10.to の末尾）などを検索"),
        ("  絞り込み", "見出しの ▼ から「ステータス」「投稿済み」「紹介文」で絞り込み"),
        ("", ""),
        ("注意", ""),
        ("  このファイルは自動で作り直されます", "Excel上で書き換えても元データには反映されません"),
        ("  Threadsリンク", "済＝SNSボタンで作ったThreads用リンクを返信に使う（SNS別レポートに載る）。未＝APIの通常リンク"),
        ("  投稿時間","商品投稿は 12:30 と 18:30 の1日2枠（投稿済みの行は実際に投稿した時刻）"),
        ("  予定日の変更", "Claude Codeに依頼するか、py drafts.py date <商品コード> <YYYY-MM-DD> [HH:MM]"),
        ("  更新されないとき", "Excelを閉じてから py sheet.py を実行"),
    ]
    for r, (label, value) in enumerate(lines, 1):
        ws.cell(row=r, column=1, value=label).font = Font(bold=not label.startswith("  ") and bool(label))
        ws.cell(row=r, column=2, value=value)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 70


def export_sheet_quietly() -> str | None:
    """他の処理のついでに管理表を更新する。失敗しても元の処理は止めない"""
    try:
        return export_sheet()
    except Exception as e:
        print(f"   ⚠️ 管理表を更新できませんでした: {e}")
        return None


if __name__ == "__main__":
    if platform.system() == "Windows":
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    print(f"✅ 管理表を更新しました: {export_sheet()}")
